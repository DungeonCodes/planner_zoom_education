"""Validar semanários e derivar DOCX exclusivamente de Markdown.

Dependências: python-docx, beautifulsoup4 e Markdown (requirements-semanarios.txt).
O QA visual é um gate separado: exportação não significa aprovação visual.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import json
import os
from pathlib import Path
import re
from zipfile import ZipFile

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
import markdown

from semanarios_paths import REPO_ROOT, SEMANARIOS_ROOT, MD_ROOT, DOCX_ROOT, docx_path, semanarios


def html(md: str) -> BeautifulSoup:
    # Markdown aceita títulos após tabelas; Python-Markdown precisa da linha
    # em branco para não absorver o título como uma célula. Normalização de
    # sintaxe para leitura/exportação, sem reescrever o arquivo fonte.
    md = re.sub(r'(?m)(^\|[^\n]*\n)(?=#{1,6} )', r'\1\n', md)
    return BeautifulSoup(markdown.markdown(md, extensions=['tables', 'sane_lists']), 'html.parser')


def words(text: str) -> list[str]:
    return re.findall(r'\w+', text.casefold())


def body_text(path: Path) -> str:
    with ZipFile(path) as z:
        assert z.testzip() is None, f'Pacote inválido: {path}'
        xml = etree.fromstring(z.read('word/document.xml'))
    return ' '.join(xml.xpath('//w:t/text()', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}))


def synchronized(path: Path) -> bool:
    dst = docx_path(path)
    return dst.exists() and words(html(path.read_text(encoding='utf-8')).get_text(' ')) == words(body_text(dst))


def local_links(path: Path) -> list[str]:
    errors = []
    soup = html(path.read_text(encoding='utf-8'))
    for a in soup.find_all('a', href=True):
        href = a['href'].split('#')[0]
        if not href or re.match(r'^[a-z]+:', href):
            continue
        from urllib.parse import unquote
        target = (path.parent / unquote(href)).resolve()
        if not target.exists():
            errors.append(f'{path.relative_to(REPO_ROOT)}: {href}')
    return errors


def infant_calendar() -> list[dict]:
    """Cruza o Plano Mestre Markdown com a grade Markdown, sem cadastro paralelo."""
    plan = REPO_ROOT / 'docs/cronograma_pedagogico/plano_mestre_infantil_2026-10-06-a-12-15.md'
    grade_path = REPO_ROOT / 'docs/cronograma_pedagogico/grade_infantil_2026.md'
    grade = {}
    for tr in html(grade_path.read_text(encoding='utf-8')).find_all('tr'):
        cells = [c.get_text(' ', strip=True) for c in tr.find_all('td')]
        if len(cells) != 6:
            continue
        day, _, start, end, turmas, prof = cells
        wd = {'Terça-feira': 1, 'Quarta-feira': 2, 'Quinta-feira': 3}[day]
        assert (datetime.strptime(end, '%H:%M') - datetime.strptime(start, '%H:%M')).seconds // 60 == 50
        for turma in re.findall(r'Inf\. (\d[A-C])', turmas):
            grade[turma] = (wd, start+'–'+end, prof)
    rows = []
    for tr in html(plan.read_text(encoding='utf-8')).find_all('tr'):
        c = [td.get_text(' ', strip=True) for td in tr.find_all('td')]
        if len(c) != 8:
            continue
        week, turma, dt, hora, activity, objective, kind, cont = c
        turma = turma.removeprefix('Inf. ')
        d = datetime.strptime(dt, '%d/%m/%Y').date()
        assert date(2026, 10, 6) <= d <= date(2026, 12, 15), c
        assert d not in {date(2026, 10, 12), date(2026, 11, 2), date(2026, 11, 15), date(2026, 11, 20)}, c
        assert grade[turma][:2] == (d.weekday(), hora), c
        assert kind in 'ABCD' and objective and cont, c
        rows.append(dict(week=week, turma=turma, dt=dt, hora=hora, activity=activity, prof=grade[turma][2]))
    assert len(rows) == 91 and len({(r['dt'], r['hora']) for r in rows}) == 81
    assert Counter(r['turma'] for r in rows) == Counter({t: 11 if t == '4B' else 10 for t in grade})
    assert len({(r['turma'], r['dt']) for r in rows}) == len(rows)
    return rows


def validate_sources() -> dict:
    assert not list(MD_ROOT.rglob('*.docx')), 'DOCX em raiz Markdown'
    assert not list(DOCX_ROOT.rglob('*.md')), 'Markdown em raiz DOCX'
    # Semanários convertidos também usam a raiz comum; originais recebidos ficam em data.
    residual = [p for base in [REPO_ROOT/'docs', REPO_ROOT/'outputs'] for p in base.rglob('semanario*')
                if p.suffix in ['.md', '.docx'] and not p.is_relative_to(MD_ROOT) and not p.is_relative_to(DOCX_ROOT)]
    assert not residual, f'Semanários fora das raízes: {residual}'
    rows = infant_calendar()
    found = []
    errors = []
    for p in semanarios():
        errors.extend(local_links(p))
        text = p.read_text(encoding='utf-8')
        assert text.strip() and '\ufffd' not in text, p
        if p.parent.name not in ['2026-10', '2026-11', '2026-12'] or 'infantil' not in p.parts:
            continue
        assert 'semana-a' not in p.name and 'semana-b' not in p.name
        assert '50 minutos' in text and 'Resumo do Encontro' in text, p
        assert 'realocar a sequência antes de avançar' not in text, p
        for tr in html(text).find_all('tr'):
            c = [td.get_text(' ', strip=True) for td in tr.find_all('td')]
            if len(c) == 5 and c[0].startswith('Inf. '):
                turma, prof, dt, hora, activity = c
                turma = turma.removeprefix('Inf. ')
            elif len(c) == 7 and re.fullmatch(r'\d[A-C]', c[0]):
                turma, prof, dh, _, activity, _, _ = c
                dt, hora = dh.split(', ')
                dt += '/2026'
            else:
                continue
            expected = [r for r in rows if r['turma'] == turma and r['dt'] == dt]
            assert len(expected) == 1, (p, c)
            assert (prof, hora, activity) == tuple(expected[0][k] for k in ['prof', 'hora', 'activity']), (p, c, expected)
            found.append((turma, dt))
    assert len(found) == 91 and len(set(found)) == 91
    assert not errors, '\n'.join(errors)
    sources = {p.relative_to(REPO_ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in semanarios()}
    return {'status': 'VALIDADO', 'encontros_turma': 91, 'horarios': 81, 'md': len(sources), 'sha256': sources}


def inline(parent, node):
    if isinstance(node, NavigableString):
        parent.add_run(re.sub(r'\s+', ' ', str(node)))
        return
    if not isinstance(node, Tag):
        return
    if node.name == 'a':
        # O rótulo visível é o mesmo do MD; a relação mantém o link clicável.
        rel = parent.part.relate_to(node.get('href', ''), RT.HYPERLINK, is_external=True)
        link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rel)
        run = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
        style = OxmlElement('w:rStyle'); style.set(qn('w:val'), 'Hyperlink'); prop.append(style)
        run.append(prop); t = OxmlElement('w:t'); t.text = node.get_text(); run.append(t); link.append(run)
        parent._p.append(link)
    elif node.name == 'br':
        parent.add_run().add_break()
    elif node.name in ['strong', 'b', 'em', 'i', 'u', 'code']:
        r = parent.add_run(node.get_text()); r.bold = node.name in ['strong', 'b']
        r.italic = node.name in ['em', 'i']; r.underline = node.name == 'u'
    else:
        for child in node.children:
            inline(parent, child)


def no_split(row):
    pr = row._tr.get_or_add_trPr()
    pr.append(OxmlElement('w:cantSplit'))


def export(path: Path):
    """Conteúdo sempre vem do MD; o DOCX anterior nunca fornece texto."""
    d = Document()
    sec = d.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
    for name, size in [('Normal', 11), ('Title', 18), ('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 11)]:
        s = d.styles[name]; s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.line_spacing = 1.08
    for name in ['Heading 1','Heading 2','Heading 3','Title']:
        d.styles[name].paragraph_format.keep_with_next = True
    # O template padrão do Word pode herdar uma borda azul em Title.
    for style in d.styles:
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    for name in ['Normal','List Bullet','List Number']:
        d.styles[name].paragraph_format.widow_control = True
    # Numeração de páginas institucional, sem conteúdo pedagógico adicional.
    fp = sec.footer.paragraphs[0]; fp.alignment = 2
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)
    soup = html(path.read_text(encoding='utf-8'))
    infant = 'infantil' in path.parts
    for node in soup.children:
        if isinstance(node, NavigableString):
            if node.strip(): d.add_paragraph(str(node).strip())
            continue
        tag = node.name
        if tag in ['h1','h2','h3','h4','h5','h6']:
            para = d.add_paragraph(style='Title' if tag=='h1' else f'Heading {min(int(tag[1])-1,3)}')
            if infant and tag=='h3' and node.get_text().startswith('Infantil '):
                para.paragraph_format.page_break_before = True
            for child in node.children: inline(para,child)
        elif tag == 'table':
            trs = node.find_all('tr'); n = len(trs[0].find_all(['td','th'],recursive=False))
            table = d.add_table(rows=0, cols=n); table.style = 'Table Grid'; table.autofit = False
            widths = [15.92/n]*n
            if n==5 and infant: widths=[1.3,1.6,2.2,2.6,8.22]
            if n==7 and infant: widths=[1.1,1.7,2.5,2.4,3.3,2.1,2.82]
            for col,w in zip(table.columns,widths): col.width=Cm(w)
            for index,tr in enumerate(trs):
                row=table.add_row(); no_split(row)
                if index==0:
                    pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:tblHeader'))
                for cell,td,w in zip(row.cells,tr.find_all(['td','th'],recursive=False),widths):
                    cell.width=Cm(w)
                    blocks=td.find_all(['p','ul','ol'],recursive=False)
                    if not blocks:
                        blocks=[td]
                    for bi,block in enumerate(blocks):
                        para=cell.paragraphs[0] if bi==0 else cell.add_paragraph()
                        para.paragraph_format.space_after=Pt(3)
                        for child in block.children: inline(para,child)
                        for run in para.runs:
                            run.font.size=Pt(9 if n>3 else 10)
                            if index==0: run.bold=True
                    if index==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)
            d.add_paragraph().paragraph_format.space_after=Pt(0)
        elif tag in ['ul','ol']:
            def add_list(list_node):
                num_id=None
                if list_node.name=='ol':
                    numbering=d.part.numbering_part.element
                    style_num=d.styles['List Number'].element.xpath('./w:pPr/w:numPr/w:numId')[0].get(qn('w:val'))
                    abstract=numbering.xpath(f'./w:num[@w:numId="{style_num}"]/w:abstractNumId')[0].get(qn('w:val'))
                    num_id=str(max(int(n.get(qn('w:numId'))) for n in numbering.xpath('./w:num'))+1)
                    num=OxmlElement('w:num');num.set(qn('w:numId'),num_id)
                    an=OxmlElement('w:abstractNumId');an.set(qn('w:val'),abstract);num.append(an)
                    override=OxmlElement('w:lvlOverride');override.set(qn('w:ilvl'),'0')
                    start=OxmlElement('w:startOverride');start.set(qn('w:val'),list_node.get('start','1'));override.append(start);num.append(override)
                    numbering.append(num)
                for li in list_node.find_all('li',recursive=False):
                    content=[child for child in li.children if not (isinstance(child,Tag) and child.name in ['ul','ol'])]
                    if any(str(child).strip() for child in content):
                        p=d.add_paragraph(style='List Number' if list_node.name=='ol' else 'List Bullet')
                        if num_id:
                            pr=p._p.get_or_add_pPr();np=OxmlElement('w:numPr')
                            level=OxmlElement('w:ilvl');level.set(qn('w:val'),'0');np.append(level)
                            ni=OxmlElement('w:numId');ni.set(qn('w:val'),num_id);np.append(ni);pr.append(np)
                        for child in content: inline(p,child)
                    for nested in li.find_all(['ul','ol'],recursive=False): add_list(nested)
            add_list(node)
        elif tag=='blockquote':
            for block in node.find_all('p',recursive=False):
                p=d.add_paragraph();p.paragraph_format.left_indent=Cm(.5)
                for child in block.children: inline(p,child)
        elif tag!='hr':
            p=d.add_paragraph()
            for child in node.children: inline(p,child)
            if re.match(r'^(segunda|terça|quarta|quinta|sexta)(-feira)?\s+\d',p.text,re.I):
                p.paragraph_format.keep_with_next=True
    target=docx_path(path);target.parent.mkdir(parents=True,exist_ok=True);d.save(target)
    assert synchronized(path), f'Conteúdo divergente após exportação: {path}'
    return target


def matrix():
    lines=['# Matriz de validação Markdown × DOCX do acervo','',
           'Markdown é a fonte de verdade. SINCRONIZADO indica igualdade lexical do texto visível, pacote DOCX íntegro e caminhos corretos. QA visual dos derivados desta execução é registrado no relatório de 06/10; pares apenas migrados conservam os bytes e o QA preexistente, sem alegar nova inspeção.', '',
           '| Segmento | Professor | Semana | Markdown | DOCX | Status |',
           '| --- | --- | --- | --- | --- | --- |']
    for p in semanarios():
        segment=p.relative_to(MD_ROOT).parts[0]
        if 'rafael' in p.name: prof='Rafael Martins'
        elif 'ricardo' in p.name: prof='Ricardo Palhares'
        elif segment=='infantil': prof='Thayane (histórico)' if p.parent.name=='2026-08' else 'Professoras da grade'
        else: prof='Tayná (histórico recebido)'
        week=re.search(r'2026-\d\d-\d\d.*',p.stem)[0]
        m=p.relative_to(SEMANARIOS_ROOT).as_posix();doc=docx_path(p).relative_to(SEMANARIOS_ROOT).as_posix()
        status='SINCRONIZADO' if synchronized(p) else 'PENDENTE'
        lines.append(f'| {segment} | {prof} | {week} | [{p.name}]({m}) | [{docx_path(p).name}]({doc}) | {status} |')
    lines+=['','Os históricos recebidos em Word permanecem preservados como fontes em `data/semanarios_prof_anterior/`; seus novos derivados são produzidos das versões Markdown, sem usar os originais para decisões pedagógicas. Nos históricos do Fundamental I, o ano 2026 no nome continua sendo convenção arquivística, conforme o README de origem.']
    (SEMANARIOS_ROOT/'matriz-validacao-md-docx.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action',choices=['validate','export','matrix'])
    parser.add_argument('--segmento',choices=['infantil','fundamental-1'])
    parser.add_argument('--only',help='Exportar somente nomes que contenham este texto')
    parser.add_argument('--out-report',type=Path)
    parser.add_argument('--missing-only',action='store_true')
    args=parser.parse_args()
    report=validate_sources()  # Gate de TODO o Markdown antes de qualquer DOCX.
    if args.out_report:
        args.out_report.parent.mkdir(parents=True,exist_ok=True)
        args.out_report.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    if args.action=='export':
        paths=semanarios(args.segmento)
        for p in paths:
            if args.only and args.only not in p.name: continue
            if args.missing_only and docx_path(p).exists(): continue
            print(export(p).relative_to(REPO_ROOT))
    elif args.action=='matrix': matrix()
    else: print(json.dumps({k:v for k,v in report.items() if k!='sha256'},ensure_ascii=False))


if __name__=='__main__': main()
